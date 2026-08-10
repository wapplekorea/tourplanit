from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


OUT = Path(__file__).resolve().parents[1] / "documents" / "TourPlanIt_제품소개서_데모가이드_2026-08.docx"
PDF_OUT = OUT.with_suffix(".pdf")
KOREAN_FONT_FILE = Path("/System/Library/AssetsV2/com_apple_MobileAsset_Font8/7a0b5c0f3c1d41c4c52a33343496c9c65ad52c50.asset/AssetData/NanumGothic.ttc")

NAVY = "163B68"
BLUE = "2E74B5"
INK = "172033"
MUTED = "61708A"
LIGHT = "EAF1F8"
PALE = "F5F8FC"
GOLD = "C68A1D"
WHITE = "FFFFFF"


KR_FONT = "Arial Unicode MS"


def set_font(run, size=11, color=INK, bold=None, italic=None):
    run.font.name = KR_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=130, bottom=90, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def para(doc, text="", size=11, color=INK, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(text)
    set_font(r, 10.7)
    return p


def heading(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: BLUE, 2: BLUE, 3: NAVY}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    p.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, sizes[level], colors[level], bold=True)
    return p


def add_header_footer(section):
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    r = header.add_run("TourPlanIt | 관광 데이터 기반 여행상품 기획 워크스페이스")
    set_font(r, 8.5, MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Wapple Studio · 제품소개서 및 데모 가이드 · 2026.08")
    set_font(r, 8.5, MUTED)


def add_kpi_strip(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [3120, 3120, 3120])
    items = [
        ("관광 데이터", "기획 재료로 탐색·선정"),
        ("AI 초안", "자동 확정 대신 검토 가능한 구조화"),
        ("연결 산출물", "기획 · 일정 · 예산 · 홍보"),
    ]
    for cell, (title, detail) in zip(table.rows[0].cells, items):
        shade(cell, PALE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        set_font(r, 10.5, NAVY, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(detail)
        set_font(r2, 9.2, MUTED)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_font(r, 9.5, NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(value)
            set_font(r, 9.3, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def callout(doc, label, text, fill="F4F6F9", color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label + "  ")
    set_font(r, 10.2, color, bold=True)
    r2 = p.add_run(text)
    set_font(r2, 10.2, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    add_header_footer(section)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = KR_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), KR_FONT)
    normal.font.size = Pt(11)

    para(doc, "2026 관광데이터 활용 공모전 · 웹·앱 구현 부문", 11, GOLD, bold=True, after=20, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "TourPlanIt", 28, NAVY, bold=True, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "관광 데이터 기반 여행상품 기획 워크스페이스", 16, INK, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "제품소개서 · 데모 가이드", 11, MUTED, after=34, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_kpi_strip(doc)
    para(doc, "", after=22)
    callout(doc, "핵심 제안", "관광자원 탐색부터 상품 기획안, 일정표, 참고 예산과 홍보 초안까지 하나의 검토 흐름으로 연결합니다.", "EAF1F8")
    para(doc, "제출자  Wapple Studio  |  기준일  2026년 8월", 9.5, MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()
    heading(doc, "1. 서비스 개요")
    para(doc, "TourPlanIt은 여행사가 목적지, 기간, 테마, 예산과 고객 조건을 입력하면 관광 데이터를 바탕으로 여행상품 기획 초안을 만드는 워크스페이스입니다. 담당자는 생성된 결과를 검토·수정한 뒤 일정, 참고 예산, 홍보 문구로 이어서 활용합니다.")
    callout(doc, "문제", "관광지 탐색, 일정 구성, 비용 정리와 홍보 문구 작성이 여러 문서와 도구에 나뉘어 반복됩니다.", "FFF8E8", "7A5A00")
    heading(doc, "2. 핵심 사용자와 사용 장면")
    add_table(doc, ["사용자", "주요 장면", "TourPlanIt이 돕는 일"], [
        ("상품기획자", "신규 목적지·테마 상품 발굴", "관광자원 후보와 상품 방향을 한 번에 초안화"),
        ("여행 OP", "출발 전 일정과 운영 조건 검토", "일자별 구성·식사·숙박·이동 누락을 검토"),
        ("지역 관광사업자", "지역 자원을 체험 상품으로 구성", "관광 데이터를 지역형 상품 스토리로 전환"),
        ("영업 담당자", "검토된 상품을 고객에게 제안", "기획 결과를 견적·블로그·카카오 문구로 연결"),
    ], [2000, 3100, 4260])
    heading(doc, "3. 사용자 흐름")
    add_table(doc, ["단계", "담당자 행동", "산출물"], [
        ("01. 조건 입력", "지역·기간·테마·대상·예산을 정리", "상품 기획 조건"),
        ("02. 데이터 탐색", "관광 데이터 후보를 확인·선정", "기획 근거와 관광 후보"),
        ("03. AI 초안", "기획안과 일자별 일정을 생성", "수정 가능한 구조화 초안"),
        ("04. 실무 검토", "숙박·식사·이동·참고 예산을 점검", "검토된 일정·참고 예산"),
        ("05. 활용·공유", "홍보 문구를 만들고 백업·공유", "고객/내부 협업용 결과물"),
    ], [1500, 4700, 3160])

    doc.add_page_break()
    heading(doc, "4. 기능 구성")
    add_table(doc, ["기능", "핵심 동작", "실무 원칙"], [
        ("관광 데이터 탐색", "목적지와 관심사에 맞는 관광 후보를 기획 재료로 조회", "관광지 기본 정보·테마·위치를 참고하고 최종 운영정보는 재확인"),
        ("AI 기획 초안", "상품 콘셉트, 일자별 일정, 소개 문구의 구조화 초안 생성", "AI는 자동 확정하지 않으며 담당자가 수정 후 사용"),
        ("일정·참고 예산", "기획서에서 일정표와 항목별 예산을 함께 확인", "가격은 확정 견적이 아닌 참고 예산 범위로 표시"),
        ("홍보 전환", "블로그·카카오·카드뉴스용 문구 초안 생성", "최종 발행 전 채널별 표현과 사실관계 검토"),
        ("보관·공유", "브라우저 이력, JSON 백업, 공유 링크, 인쇄", "개인정보·계약조건·원가는 공유 문서에 넣지 않음"),
    ], [1900, 3800, 3660])
    heading(doc, "5. 관광 데이터 활용 및 보호 원칙")
    bullet(doc, "관광 데이터는 목적지별 관광자원 후보를 찾고 상품 기획의 근거를 만드는 데 사용합니다.")
    bullet(doc, "관광지 운영시간, 예약 가능 여부, 최신 요금과 현장 상황은 최종 확정 전 담당자가 별도로 확인합니다.")
    bullet(doc, "관광 데이터와 개인 고객 정보는 분리하며, 제출·데모에는 비식별 예시만 사용합니다.")
    bullet(doc, "관광공사 API 키와 AI API 키는 브라우저가 아니라 서버 함수에서만 처리합니다.")
    callout(doc, "현재 MVP 범위", "관광 데이터 후보와 AI 결과는 검토용 초안입니다. 자동 예약·자동 발권·자동 확정 기능은 제공하지 않습니다.", "FFF4F4", "9B1C1C")
    heading(doc, "6. 차별점")
    para(doc, "TourPlanIt의 차별점은 단순한 여행 일정 생성이 아니라, 관광 데이터와 상품기획·일정·참고 예산·홍보 초안을 하나의 검토 가능한 결과로 연결하는 데 있습니다. 담당자가 생성에 쓰던 시간을 줄이고 동선·현장성·예산 검토에 집중하도록 돕습니다.")

    doc.add_page_break()
    heading(doc, "7. 3분 데모 진행안")
    add_table(doc, ["시간", "화면", "시연 포인트"], [
        ("0:00–0:25", "기획 시작", "여행상품 기획 업무를 한 화면에서 시작한다는 문제 정의"),
        ("0:25–1:05", "조건 입력·관광 후보", "지역·기간·테마 조건과 관광 데이터가 기획 재료로 연결됨"),
        ("1:05–1:50", "기획·일정 편집", "AI가 초안을 만들고 담당자가 실무 기준으로 검토·수정함"),
        ("1:50–2:30", "참고 예산·홍보", "같은 기획 결과가 예산과 고객용 문구로 이어짐"),
        ("2:30–3:00", "점검·공유·백업", "검토 경고 확인 후 개인정보 없는 문서만 공유·보관"),
    ], [1500, 2500, 5360])
    heading(doc, "8. 제출 전 최종 점검")
    for item in [
        "관광 데이터 탐색 → 기획안 → 일정 수정 → 참고 예산 → 홍보 문구 → 백업/공유를 한 건으로 시연한다.",
        "데스크톱, 태블릿, 폴드 와이드, 390px 모바일에서 긴 제목·버튼·가로 스크롤을 점검한다.",
        "AI·관광 데이터 오류 시 사용자가 재시도할 수 있는 안내가 보이는지 점검한다.",
        "공유 URL에는 개인정보, 계약조건, 원가가 포함되지 않는지 확인한다.",
        "Netlify 서버 환경변수와 Function 정상·오류 응답을 배포 직후 확인한다.",
    ]:
        bullet(doc, item)
    heading(doc, "9. 최종 배포 원칙")
    para(doc, "제출 전 내부 QA와 문서 확정을 먼저 완료한 뒤 Netlify에 한 번 배포합니다. 운영 환경에는 ANTHROPIC_API_KEY와 KTO_API_KEY만 서버 환경변수로 등록하며, VITE_ 접두사의 비밀키는 사용하지 않습니다.")
    callout(doc, "배포 후 확인", "관광 데이터 조회 · AI 기획 초안 · 공유 링크 · 인쇄 · 모바일 화면을 운영 URL에서 한 번씩 재검증합니다.", "EAF1F8")

    doc.save(OUT)
    build_pdf()
    print(OUT)


def build_pdf():
    """Create the submission PDF separately so the Korean layout is stable on every viewer."""
    pdfmetrics.registerFont(TTFont("TourNanum", str(KOREAN_FONT_FILE), subfontIndex=0))
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "TourBase", parent=styles["BodyText"], fontName="TourNanum", fontSize=9.6,
        leading=15, textColor=colors.HexColor("#172033"), spaceAfter=5,
    )
    title = ParagraphStyle("TourTitle", parent=base, fontSize=28, leading=34, alignment=TA_CENTER, textColor=colors.HexColor("#163B68"), spaceAfter=7)
    subtitle = ParagraphStyle("TourSubtitle", parent=base, fontSize=15, leading=21, alignment=TA_CENTER, textColor=colors.HexColor("#172033"), spaceAfter=10)
    eyebrow = ParagraphStyle("TourEyebrow", parent=base, fontSize=10.5, leading=15, alignment=TA_CENTER, textColor=colors.HexColor("#C68A1D"), spaceAfter=22)
    heading_style = ParagraphStyle("TourHeading", parent=base, fontSize=15, leading=21, textColor=colors.HexColor("#2E74B5"), spaceBefore=13, spaceAfter=6)
    small = ParagraphStyle("TourSmall", parent=base, fontSize=8.7, leading=12, textColor=colors.HexColor("#61708A"))

    def P(value, style=base):
        return Paragraph(value.replace("\n", "<br/>"), style)

    def grid(headers, rows, widths):
        data = [[P(value, ParagraphStyle("head", parent=small, textColor=colors.HexColor("#163B68"), fontSize=8.7, leading=11)) for value in headers]]
        data += [[P(value, ParagraphStyle("cell", parent=base, fontSize=8.4, leading=11.5)) for value in row] for row in rows]
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1F8")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CED9E6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        return table

    def note(label, text, fill="#EAF1F8"):
        table = Table([[P(f"<b>{label}</b>  {text}")]], colWidths=[176 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(fill)),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD9E9")),
            ("LEFTPADDING", (0, 0), (-1, -1), 9),
            ("RIGHTPADDING", (0, 0), (-1, -1), 9),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ]))
        return table

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("TourNanum", 8)
        canvas.setFillColor(colors.HexColor("#61708A"))
        canvas.drawCentredString(A4[0] / 2, 11 * mm, "Wapple Studio · 제품소개서 및 데모 가이드 · 2026.08")
        canvas.drawRightString(A4[0] - 17 * mm, A4[1] - 12 * mm, f"TourPlanIt | {document.page}")
        canvas.restoreState()

    story = [
        Spacer(1, 30 * mm),
        P("2026 관광데이터 활용 공모전 · 웹·앱 구현 부문", eyebrow),
        P("TourPlanIt", title),
        P("관광 데이터 기반 여행상품 기획 워크스페이스", subtitle),
        P("제품소개서 · 데모 가이드", ParagraphStyle("coverSub", parent=small, alignment=TA_CENTER, fontSize=11, leading=16, spaceAfter=26)),
        note("핵심 제안", "관광자원 탐색부터 상품 기획안, 일정표, 참고 예산과 홍보 초안까지 하나의 검토 흐름으로 연결합니다."),
        Spacer(1, 9 * mm),
        P("제출자  Wapple Studio  |  기준일  2026년 8월", ParagraphStyle("coverFooter", parent=small, alignment=TA_CENTER)),
        PageBreak(),
        P("1. 서비스 개요", heading_style),
        P("TourPlanIt은 여행사가 목적지, 기간, 테마, 예산과 고객 조건을 입력하면 관광 데이터를 바탕으로 여행상품 기획 초안을 만드는 워크스페이스입니다. 담당자는 생성된 결과를 검토·수정한 뒤 일정, 참고 예산, 홍보 문구로 이어서 활용합니다."),
        note("문제", "관광지 탐색, 일정 구성, 비용 정리와 홍보 문구 작성이 여러 문서와 도구에 나뉘어 반복됩니다.", "#FFF8E8"),
        P("2. 핵심 사용자와 사용 장면", heading_style),
        grid(["사용자", "주요 장면", "TourPlanIt이 돕는 일"], [
            ("상품기획자", "신규 목적지·테마 상품 발굴", "관광자원 후보와 상품 방향을 한 번에 초안화"),
            ("여행 OP", "출발 전 일정과 운영 조건 검토", "일자별 구성·식사·숙박·이동 누락을 검토"),
            ("지역 관광사업자", "지역 자원을 체험 상품으로 구성", "관광 데이터를 지역형 상품 스토리로 전환"),
            ("영업 담당자", "검토된 상품을 고객에게 제안", "기획 결과를 견적·블로그·카카오 문구로 연결"),
        ], [35 * mm, 57 * mm, 84 * mm]),
        P("3. 사용자 흐름", heading_style),
        grid(["단계", "담당자 행동", "산출물"], [
            ("01. 조건 입력", "지역·기간·테마·대상·예산을 정리", "상품 기획 조건"),
            ("02. 데이터 탐색", "관광 데이터 후보를 확인·선정", "기획 근거와 관광 후보"),
            ("03. AI 초안", "기획안과 일자별 일정을 생성", "수정 가능한 구조화 초안"),
            ("04. 실무 검토", "숙박·식사·이동·참고 예산을 점검", "검토된 일정·참고 예산"),
            ("05. 활용·공유", "홍보 문구를 만들고 백업·공유", "고객/내부 협업용 결과물"),
        ], [28 * mm, 86 * mm, 62 * mm]),
        PageBreak(),
        P("4. 기능 구성", heading_style),
        grid(["기능", "핵심 동작", "실무 원칙"], [
            ("관광 데이터 탐색", "목적지·관심사에 맞는 관광 후보를 기획 재료로 조회", "운영시간·요금·현장 상황은 최종 확정 전 재확인"),
            ("AI 기획 초안", "콘셉트·일자별 일정·소개 문구의 구조화 초안 생성", "AI는 자동 확정하지 않고 담당자가 수정 후 사용"),
            ("일정·참고 예산", "기획서에서 일정표와 항목별 예산을 함께 확인", "가격은 확정 견적이 아닌 참고 예산 범위"),
            ("홍보 전환", "블로그·카카오·카드뉴스용 문구 초안 생성", "최종 발행 전 채널별 표현과 사실관계 검토"),
            ("보관·공유", "브라우저 이력, JSON 백업, 공유 링크, 인쇄", "개인정보·계약조건·원가는 공유 문서에 넣지 않음"),
        ], [35 * mm, 71 * mm, 70 * mm]),
        P("5. 관광 데이터 활용 및 보호 원칙", heading_style),
        *[P(f"• {line}") for line in [
            "관광 데이터는 목적지별 관광자원 후보를 찾고 상품 기획의 근거를 만드는 데 사용합니다.",
            "관광지 운영시간, 예약 가능 여부, 최신 요금과 현장 상황은 최종 확정 전 담당자가 별도로 확인합니다.",
            "관광 데이터와 개인 고객 정보는 분리하며, 제출·데모에는 비식별 예시만 사용합니다.",
            "관광공사 API 키와 AI API 키는 브라우저가 아니라 서버 함수에서만 처리합니다.",
        ]],
        note("현재 MVP 범위", "관광 데이터 후보와 AI 결과는 검토용 초안입니다. 자동 예약·자동 발권·자동 확정 기능은 제공하지 않습니다.", "#FFF4F4"),
        P("6. 차별점", heading_style),
        P("TourPlanIt은 단순한 여행 일정 생성이 아니라, 관광 데이터와 상품기획·일정·참고 예산·홍보 초안을 하나의 검토 가능한 결과로 연결합니다. 담당자가 생성에 쓰던 시간을 줄이고 동선·현장성·예산 검토에 집중하도록 돕습니다."),
        PageBreak(),
        P("7. 3분 데모 진행안", heading_style),
        grid(["시간", "화면", "시연 포인트"], [
            ("0:00–0:25", "기획 시작", "여행상품 기획 업무를 한 화면에서 시작한다는 문제 정의"),
            ("0:25–1:05", "조건 입력·관광 후보", "지역·기간·테마 조건과 관광 데이터가 기획 재료로 연결됨"),
            ("1:05–1:50", "기획·일정 편집", "AI가 초안을 만들고 담당자가 실무 기준으로 검토·수정함"),
            ("1:50–2:30", "참고 예산·홍보", "같은 기획 결과가 예산과 고객용 문구로 이어짐"),
            ("2:30–3:00", "점검·공유·백업", "검토 경고 확인 후 개인정보 없는 문서만 공유·보관"),
        ], [29 * mm, 47 * mm, 100 * mm]),
        P("8. 제출 전 최종 점검", heading_style),
        *[P(f"• {line}") for line in [
            "관광 데이터 탐색 → 기획안 → 일정 수정 → 참고 예산 → 홍보 문구 → 백업/공유를 한 건으로 시연합니다.",
            "데스크톱, 태블릿, 폴드 와이드, 390px 모바일에서 긴 제목·버튼·가로 스크롤을 점검합니다.",
            "AI·관광 데이터 오류 시 사용자가 재시도할 수 있는 안내가 보이는지 점검합니다.",
            "공유 URL에는 개인정보, 계약조건, 원가가 포함되지 않는지 확인합니다.",
            "Netlify 서버 환경변수와 Function 정상·오류 응답을 배포 직후 확인합니다.",
        ]],
        P("9. 최종 배포 원칙", heading_style),
        P("제출 전 내부 QA와 문서 확정을 먼저 완료한 뒤 Netlify에 한 번 배포합니다. 운영 환경에는 ANTHROPIC_API_KEY와 KTO_API_KEY만 서버 환경변수로 등록하며, VITE_ 접두사의 비밀키는 사용하지 않습니다."),
        note("배포 후 확인", "관광 데이터 조회 · AI 기획 초안 · 공유 링크 · 인쇄 · 모바일 화면을 운영 URL에서 한 번씩 재검증합니다."),
    ]
    SimpleDocTemplate(
        str(PDF_OUT), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm,
        topMargin=19 * mm, bottomMargin=18 * mm, title="TourPlanIt 제품소개서 및 데모 가이드",
    ).build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    build()
